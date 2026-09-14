import io

import pymupdf
from PIL import Image


def _make_test_image() -> bytes:
    img = Image.new("RGB", (300, 200), color=(200, 50, 50))
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def test_image_to_pdf(client):
    response = client.post(
        "/api/v1/convert/image-to-pdf",
        files=[("images", ("photo.png", _make_test_image(), "image/png"))],
        data={"orientation": "auto", "margin": "small", "pageSize": "a4"},
    )
    assert response.status_code == 200
    doc = pymupdf.open(stream=response.content, filetype="pdf")
    assert doc.page_count == 1


def test_word_to_pdf(client):
    from docx import Document

    document = Document()
    document.add_paragraph("Hello from a Word document.")
    buffer = io.BytesIO()
    document.save(buffer)

    response = client.post(
        "/api/v1/convert/word-to-pdf",
        files={
            "file": (
                "test.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    result_doc = pymupdf.open(stream=response.content, filetype="pdf")
    assert result_doc.page_count >= 1
    assert "Hello from a Word document" in result_doc[0].get_text()


def test_excel_to_pdf(client):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "NSSF Test Sheet"
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = client.post(
        "/api/v1/convert/excel-to-pdf",
        files={
            "file": (
                "test.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    result_doc = pymupdf.open(stream=response.content, filetype="pdf")
    assert result_doc.page_count >= 1


def test_powerpoint_to_pdf(client):
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "NSSF Slide"
    buffer = io.BytesIO()
    presentation.save(buffer)

    response = client.post(
        "/api/v1/convert/powerpoint-to-pdf",
        files={
            "file": (
                "test.pptx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    assert response.status_code == 200
    result_doc = pymupdf.open(stream=response.content, filetype="pdf")
    assert result_doc.page_count >= 1


def test_html_to_pdf_from_raw_html(client):
    response = client.post(
        "/api/v1/convert/html-to-pdf",
        data={
            "inputMode": "html",
            "html": "<html><body><h1>NSSF Report</h1></body></html>",
            "width": "800",
            "printBackground": "true",
            "margin": "small",
        },
    )
    assert response.status_code == 200
    doc = pymupdf.open(stream=response.content, filetype="pdf")
    assert doc.page_count >= 1
    assert "NSSF Report" in doc[0].get_text()
