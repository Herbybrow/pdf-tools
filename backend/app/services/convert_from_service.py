import io
import logging
import time

import pdfplumber
import pymupdf
from openpyxl import Workbook
from pdf2docx import Converter
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from app.core.engines import find_ghostscript
from app.core.page_ranges import parse_page_ranges
from app.core.subprocess_utils import run_cli
from app.core.tempfiles import temp_workspace
from app.services.optimize_service import ocr_pdf

logger = logging.getLogger("nssf.conversion")
_EMU_PER_POINT = 12700


def _order_blocks_by_reading_order(blocks: list[tuple], page_width: float) -> list[tuple]:
    if len(blocks) < 2:
        return blocks

    candidate_splits = [page_width * fraction for fraction in (0.45, 0.48, 0.5, 0.52, 0.55)]
    gutter_x = None
    for split in candidate_splits:
        if not any(b[0] < split < b[2] for b in blocks):
            left_count = sum(1 for b in blocks if (b[0] + b[2]) / 2 < split)
            right_count = len(blocks) - left_count
            if left_count >= 2 and right_count >= 2:
                gutter_x = split
                break

    if gutter_x is None:
        return sorted(blocks, key=lambda b: (round(b[1] / 10), b[0]))

    left_column = sorted((b for b in blocks if (b[0] + b[2]) / 2 < gutter_x), key=lambda b: b[1])
    right_column = sorted((b for b in blocks if (b[0] + b[2]) / 2 >= gutter_x), key=lambda b: b[1])
    return left_column + right_column


def _infer_alignment(x0: float, x1: float, column_left: float, column_right: float) -> PP_ALIGN:
    column_width = column_right - column_left
    block_width = x1 - x0
    if block_width > column_width * 0.75:
        return PP_ALIGN.LEFT
    left_margin = x0 - column_left
    right_margin = column_right - x1
    if abs(left_margin - right_margin) < 12:
        return PP_ALIGN.CENTER
    if right_margin < left_margin - 12 and right_margin < 12:
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def pdf_to_word(data: bytes, ocr_scanned: bool = False) -> bytes:
    t0 = time.perf_counter()
    source = ocr_pdf(data, "eng") if ocr_scanned else data
    with temp_workspace() as workspace:
        input_path = workspace / "input.pdf"
        output_path = workspace / "output.docx"
        input_path.write_bytes(source)
        converter = Converter(str(input_path))
        try:
            converter.convert(str(output_path))
        finally:
            converter.close()
        if not output_path.exists():
            raise RuntimeError("Conversion did not produce a Word document.")
        out = output_path.read_bytes()
        logger.info(f"[TIMING] pdf_to_word executed in {time.perf_counter() - t0:.2f}s")
        return out


def pdf_to_powerpoint(data: bytes) -> bytes:
    t0 = time.perf_counter()
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        if doc.page_count == 0:
            raise ValueError("This PDF has no pages.")
        presentation = Presentation()
        first_rect = doc[0].rect
        presentation.slide_width = Emu(int(first_rect.width * _EMU_PER_POINT))
        presentation.slide_height = Emu(int(first_rect.height * _EMU_PER_POINT))
        blank_layout = presentation.slide_layouts[6]

        for page in doc:
            slide = presentation.slides.add_slide(blank_layout)
            page_width = page.rect.width
            text_blocks = [b for b in page.get_text("blocks") if b[4].strip()]
            ordered_blocks = _order_blocks_by_reading_order(text_blocks, page_width)

            for x0, y0, x1, y1, text, *_rest in ordered_blocks:
                text = text.strip()
                textbox = slide.shapes.add_textbox(
                    Emu(int(x0 * _EMU_PER_POINT)),
                    Emu(int(y0 * _EMU_PER_POINT)),
                    Emu(int(max(x1 - x0, 10) * _EMU_PER_POINT)),
                    Emu(int(max(y1 - y0, 10) * _EMU_PER_POINT)),
                )
                text_frame = textbox.text_frame
                text_frame.word_wrap = True
                text_frame.text = text
                alignment = _infer_alignment(x0, x1, 0, page_width)
                for paragraph in text_frame.paragraphs:
                    paragraph.alignment = alignment
                    for run in paragraph.runs:
                        run.font.size = Pt(11)

            seen_xrefs: set[int] = set()
            for image_info in page.get_images(full=True):
                xref = image_info[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    rects = page.get_image_rects(xref)
                except ValueError:
                    rects = []
                if not rects:
                    continue
                base_image = doc.extract_image(xref)
                image_stream = io.BytesIO(base_image["image"])
                rect = rects[0]
                slide.shapes.add_picture(
                    image_stream,
                    Emu(int(rect.x0 * _EMU_PER_POINT)),
                    Emu(int(rect.y0 * _EMU_PER_POINT)),
                    Emu(int(max(rect.x1 - rect.x0, 1) * _EMU_PER_POINT)),
                    Emu(int(max(rect.y1 - rect.y0, 1) * _EMU_PER_POINT)),
                )

        buffer = io.BytesIO()
        presentation.save(buffer)
        logger.info(f"[TIMING] pdf_to_powerpoint executed in {time.perf_counter() - t0:.2f}s")
        return buffer.getvalue()
    finally:
        doc.close()


_TEXT_TABLE_STRATEGY = {
    "vertical_strategy": "text",
    "horizontal_strategy": "text",
    "intersection_tolerance": 5,
}


def _looks_like_real_table(tables: list[list[list[str | None]]]) -> bool:
    return any(len(row) >= 2 for table in tables for row in table)


def pdf_to_excel(data: bytes, pages_spec: str | None = None, ocr_scanned: bool = False) -> bytes:
    t0 = time.perf_counter()
    source = ocr_pdf(data, "eng") if ocr_scanned else data
    workbook = Workbook()
    workbook.remove(workbook.active)
    any_table_found = False

    with pdfplumber.open(io.BytesIO(source)) as pdf:
        total_pages = len(pdf.pages)
        page_indices = parse_page_ranges(pages_spec, total_pages) if pages_spec else list(range(total_pages))

        for page_index in page_indices:
            page = pdf.pages[page_index]
            tables = page.extract_tables()
            if not tables or not _looks_like_real_table(tables):
                text_tables = page.extract_tables(table_settings=_TEXT_TABLE_STRATEGY)
                tables = text_tables if text_tables and _looks_like_real_table(text_tables) else None
            if not tables:
                continue
            any_table_found = True
            sheet = workbook.create_sheet(title=f"Page {page_index + 1}"[:31])
            current_row = 1
            for table in tables:
                for row in table:
                    for col_index, cell_value in enumerate(row, start=1):
                        sheet.cell(row=current_row, column=col_index, value=cell_value)
                    current_row += 1
                current_row += 1

        if not any_table_found:
            all_lines: list[str] = []
            for page_index in page_indices:
                text = pdf.pages[page_index].extract_text() or ""
                all_lines.extend(text.splitlines())

            if not any(line.strip() for line in all_lines):
                if ocr_scanned:
                    raise ValueError(
                        "No text or tables could be found on the selected pages, even after OCR. "
                        "The scan may be too low-quality to read."
                    )
                raise ValueError(
                    "This PDF has no selectable text on the selected pages (it looks like a scanned "
                    'image). Enable "Scanned document (use OCR)" and try again.'
                )

            sheet = workbook.create_sheet(title="Extracted Text")
            for row_num, line in enumerate(all_lines, start=1):
                sheet.cell(row=row_num, column=1, value=line)

    buffer = io.BytesIO()
    workbook.save(buffer)
    logger.info(f"[TIMING] pdf_to_excel executed in {time.perf_counter() - t0:.2f}s")
    return buffer.getvalue()


def pdf_to_images(
    data: bytes,
    image_format: str = "jpg",
    dpi: int = 150,
    extract_embedded_images: bool = False,
) -> dict[str, bytes]:
    t0 = time.perf_counter()
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        ext = "png" if image_format == "png" else "jpg"

        if extract_embedded_images:
            files: dict[str, bytes] = {}
            seen_xrefs: set[int] = set()
            counter = 1
            for page_index in range(doc.page_count):
                for image_info in doc.get_page_images(page_index, full=True):
                    xref = image_info[0]
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)
                    base_image = doc.extract_image(xref)
                    files[f"image_{counter}.{base_image['ext']}"] = base_image["image"]
                    counter += 1
            if not files:
                raise ValueError("No embedded images were found in this PDF.")
            logger.info(f"[TIMING] pdf_to_images (extract) executed in {time.perf_counter() - t0:.2f}s")
            return files

        zoom = dpi / 72
        matrix = pymupdf.Matrix(zoom, zoom)
        width = len(str(doc.page_count))
        files = {}
        for page_index in range(doc.page_count):
            pixmap = doc[page_index].get_pixmap(matrix=matrix)
            files[f"page_{str(page_index + 1).zfill(width)}.{ext}"] = pixmap.tobytes(ext)
        logger.info(f"[TIMING] pdf_to_images executed in {time.perf_counter() - t0:.2f}s")
        return files
    finally:
        doc.close()


_PDFA_LEVELS = {"pdfa-1b": "1", "pdfa-2b": "2", "pdfa-3b": "3"}


def pdf_to_pdfa(data: bytes, standard: str = "pdfa-2b") -> bytes:
    t0 = time.perf_counter()
    gs_path = find_ghostscript()
    if not gs_path:
        raise RuntimeError("Ghostscript is not installed.")
    pdfa_level = _PDFA_LEVELS.get(standard, "2")
    with temp_workspace() as workspace:
        input_path = workspace / "input.pdf"
        output_path = workspace / "output.pdf"
        input_path.write_bytes(data)
        run_cli(
            [
                gs_path,
                f"-dPDFA={pdfa_level}",
                "-dBATCH",
                "-dNOPAUSE",
                "-dQUIET",
                "-sColorConversionStrategy=RGB",
                "-sProcessColorModel=DeviceRGB",
                "-sDEVICE=pdfwrite",
                "-dPDFACompatibilityPolicy=1",
                f"-sOutputFile={output_path}",
                str(input_path),
            ],
        )
        if not output_path.exists():
            raise RuntimeError("Ghostscript could not produce a PDF/A output.")
        out = output_path.read_bytes()
        logger.info(f"[TIMING] pdf_to_pdfa executed in {time.perf_counter() - t0:.2f}s")
        return out


def pdf_to_markdown(data: bytes) -> bytes:
    t0 = time.perf_counter()
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        sizes: list[int] = []
        for page in doc:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        if span["text"].strip():
                            sizes.append(round(span["size"]))
        body_size = max(set(sizes), key=sizes.count) if sizes else 11

        lines_out: list[str] = []
        for page in doc:
            for block in page.get_text("dict")["blocks"]:
                block_text_parts: list[str] = []
                max_size = body_size
                for line in block.get("lines", []):
                    line_parts = [span["text"] for span in line.get("spans", []) if span["text"].strip()]
                    if not line_parts:
                        continue
                    block_text_parts.append("".join(line_parts))
                    max_size = max(max_size, *(round(span["size"]) for span in line.get("spans", [])))

                text = " ".join(block_text_parts).strip()
                if not text:
                    continue

                ratio = max_size / body_size if body_size else 1
                stripped = text.lstrip("-•* ").strip()
                if ratio >= 1.6:
                    lines_out.append(f"# {text}")
                elif ratio >= 1.35:
                    lines_out.append(f"## {text}")
                elif ratio >= 1.15:
                    lines_out.append(f"### {text}")
                elif text.startswith(("- ", "• ", "* ")):
                    lines_out.append(f"- {stripped}")
                else:
                    lines_out.append(text)
                lines_out.append("")
        out = "\n".join(lines_out).encode("utf-8")
        logger.info(f"[TIMING] pdf_to_markdown executed in {time.perf_counter() - t0:.2f}s")
        return out
    finally:
        doc.close()